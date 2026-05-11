import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from fastapi import HTTPException, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.models.dataset import Dataset
from app.models.prediction import Prediction
from app.models.report import Report
from app.models.user import User
from app.schemas.report import GenerateReportRequest


async def get_owned_prediction(
    db: AsyncSession,
    current_user: User,
    prediction_id: int,
) -> Prediction:
    result = await db.execute(
        select(Prediction)
        .join(Dataset)
        .where(
            Prediction.id == prediction_id,
            Dataset.owner_id == current_user.id,
        )
    )

    prediction = result.scalar_one_or_none()

    if not prediction:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Prediction not found",
        )

    return prediction


def ensure_reports_dir() -> None:
    os.makedirs(settings.REPORTS_DIR, exist_ok=True)


def generate_excel_report_file(
    prediction: Prediction,
) -> tuple[str, str]:
    ensure_reports_dir()

    filename = f"forecast_report_{prediction.id}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.xlsx"
    file_path = os.path.join(settings.REPORTS_DIR, filename)

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Forecast Report"

    header_fill = PatternFill(
        start_color="00D4FF",
        end_color="00D4FF",
        fill_type="solid",
    )

    dark_fill = PatternFill(
        start_color="0D1525",
        end_color="0D1525",
        fill_type="solid",
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
    )

    border = Border(
        left=Side(style="thin", color="D9E2EF"),
        right=Side(style="thin", color="D9E2EF"),
        top=Side(style="thin", color="D9E2EF"),
        bottom=Side(style="thin", color="D9E2EF"),
    )

    ws["A1"] = "ForecastIQ Demand Forecast Report"
    ws["A1"].font = Font(size=18, bold=True, color="00D4FF")

    ws["A3"] = "Prediction ID"
    ws["B3"] = prediction.id

    ws["A4"] = "Dataset ID"
    ws["B4"] = prediction.dataset_id

    ws["A5"] = "Model Type"
    ws["B5"] = prediction.model_type

    ws["A6"] = "Periods"
    ws["B6"] = prediction.periods

    ws["A7"] = "MAE"
    ws["B7"] = prediction.mae

    ws["A8"] = "RMSE"
    ws["B8"] = prediction.rmse

    ws["A9"] = "MAPE"
    ws["B9"] = prediction.mape

    headers = [
        "Date",
        "Forecast",
        "Lower Bound",
        "Upper Bound",
    ]

    start_row = 12

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_index)
        cell.value = header
        cell.fill = header_fill
        cell.font = white_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(prediction.results, start=start_row + 1):
        ws.cell(row=row_index, column=1).value = row.get("date")
        ws.cell(row=row_index, column=2).value = row.get("forecast")
        ws.cell(row=row_index, column=3).value = row.get("lower")
        ws.cell(row=row_index, column=4).value = row.get("upper")

        for col_index in range(1, 5):
            ws.cell(row=row_index, column=col_index).border = border

    for column in ["A", "B", "C", "D"]:
        ws.column_dimensions[column].width = 22

    summary = wb.create_sheet("Summary")

    summary["A1"] = "ForecastIQ Summary"
    summary["A1"].font = Font(size=18, bold=True, color="00D4FF")
    summary_data = [
        ["Metric", "Value"],
        ["Model Type", prediction.model_type],
        ["Forecast Days", prediction.periods],
        ["MAE", prediction.mae],
        ["RMSE", prediction.rmse],
        ["MAPE", f"{prediction.mape}%"],
    ]

    for row_index, row in enumerate(summary_data, start=3):
        for col_index, value in enumerate(row, start=1):
            cell = summary.cell(row=row_index, column=col_index)
            cell.value = value
            cell.border = border

            if row_index == 3:
                cell.fill = dark_fill
                cell.font = white_font

    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 24

    wb.save(file_path)

    return file_path, filename


def generate_pdf_report_file(
    prediction: Prediction,
) -> tuple[str, str]:
    ensure_reports_dir()

    filename = f"forecast_report_{prediction.id}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.pdf"
    file_path = os.path.join(settings.REPORTS_DIR, filename)

    doc = SimpleDocTemplate(
        file_path,
        pagesize=letter,
    )

    story = []

    title = Paragraph(
        "<b>ForecastIQ Demand Forecast Report</b>",
    )

    story.append(title)
    story.append(Spacer(1, 16))

    summary_table_data = [
        ["Metric", "Value"],
        ["Prediction ID", prediction.id],
        ["Dataset ID", prediction.dataset_id],
        ["Model Type", prediction.model_type],
        ["Forecast Days", prediction.periods],
        ["MAE", prediction.mae],
        ["RMSE", prediction.rmse],
        ["MAPE", f"{prediction.mape}%"],
    ]

    summary_table = Table(summary_table_data)

    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#00D4FF")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ]
        )
    )

    story.append(summary_table)
    story.append(Spacer(1, 24))

    forecast_rows = [
        ["Date", "Forecast", "Lower", "Upper"]
    ]

    for row in prediction.results[:40]:
        forecast_rows.append(
            [
                row.get("date"),
                row.get("forecast"),
                row.get("lower"),
                row.get("upper"),
            ]
        )

    forecast_table = Table(forecast_rows)

    forecast_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1525")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]
        )
    )

    story.append(forecast_table)

    doc.build(story)

    return file_path, filename


async def generate_report(
    db: AsyncSession,
    current_user: User,
    payload: GenerateReportRequest,
) -> Report:
    prediction = await get_owned_prediction(
        db,
        current_user,
        payload.prediction_id,
    )

    if payload.format == "excel":
        file_path, filename = generate_excel_report_file(prediction)
    elif payload.format == "pdf":
        file_path, filename = generate_pdf_report_file(prediction)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid report format",
        )

    report = Report(
        prediction_id=prediction.id,
        format=payload.format,
        file_path=file_path,
        filename=filename,
    )

    db.add(report)

    await db.commit()
    await db.refresh(report)

    return report
async def get_reports(
    db: AsyncSession,
    current_user: User,
):
    result = await db.execute(
        select(Report)
        .join(Prediction)
        .join(Dataset)
        .where(Dataset.owner_id == current_user.id)
        .order_by(Report.created_at.desc())
    )

    return result.scalars().all()


async def get_report_for_download(
    db: AsyncSession,
    current_user: User,
    report_id: int,
    report_format: str,
) -> Report:
    result = await db.execute(
        select(Report)
        .join(Prediction)
        .join(Dataset)
        .where(
            Report.id == report_id,
            Report.format == report_format,
            Dataset.owner_id == current_user.id,
        )
    )

    report = result.scalar_one_or_none()

    if not report:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Report not found",
        )

    if not os.path.exists(report.file_path):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Report file missing on server",
        )

    return report